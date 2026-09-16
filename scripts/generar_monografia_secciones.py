"""Genera documento Word con secciones iniciales de monografía Contalogic AI+."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

OUT = r"c:\Users\jazz2\Downloads\Monografia Contalogic AI+ - Resumen Introduccion Problema Justificacion.docx"


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for level, size in [(1, 14), (2, 13)]:
        style_name = f"Heading {level}"
        if style_name in doc.styles:
            h = doc.styles[style_name]
            h.font.name = "Times New Roman"
            h.font.size = Pt(size)
            h.font.bold = True


def add_paragraphs(doc: Document, text: str, justify: bool = True) -> None:
    for block in text.strip().split("\n\n"):
        p = doc.add_paragraph(block.strip())
        if justify:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(1.25)


def add_title_page(doc: Document) -> None:
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("MONOGRAFÍA\n\n")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Times New Roman"

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run(
        "Contalogic AI+ — Terminal Financiera Universal:\n"
        "Plataforma web para análisis financiero y asistencia tributaria\n"
        "con OCR e inteligencia artificial en el contexto colombiano"
    )
    r.font.size = Pt(13)
    r.font.name = "Times New Roman"

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ir = info.add_run(
        "Resumen · Introducción · Problema de investigación · Justificación\n\n"
        "Fecha de elaboración: septiembre de 2026"
    )
    ir.font.size = Pt(12)
    ir.font.name = "Times New Roman"

    doc.add_page_break()


def main() -> None:
    doc = Document()
    setup_styles(doc)
    add_title_page(doc)

    # --- RESUMEN ---
    doc.add_heading("RESUMEN", level=1)
    add_paragraphs(
        doc,
        """La gestión financiera y el cumplimiento de obligaciones tributarias constituyen uno de los principales retos para personas naturales, personas jurídicas y profesionales contables en Colombia. Frente a un entorno normativo complejo —regido por el Estatuto Tributario, las disposiciones de la DIAN y la Ley 1581 de 2012 sobre protección de datos personales—, muchos contribuyentes continúan apoyándose en herramientas dispersas (hojas de cálculo, archivos PDF, recordatorios manuales y consultas aisladas), lo que incrementa el riesgo de errores, retrasos y decisiones tributarias ineficientes.

El presente trabajo documenta el diseño, desarrollo y fundamentación de Contalogic AI+ — Terminal Financiera Universal, una plataforma web SaaS orientada al contexto colombiano, que integra en un mismo ecosistema digital la captura y extracción automatizada de estados financieros mediante OCR en el navegador, la gestión de ingresos y gastos por sector económico, tableros de control financiero, recordatorios de vencimientos DIAN, recomendaciones de ahorro tributario y un asistente conversacional con inteligencia artificial (LEX AI) especializado en normativa tributaria nacional.

La solución se construyó sobre una arquitectura cliente-servidor desacoplada, empleando React 19 con TypeScript, TanStack Router/Start, Vite y Tailwind CSS en el frontend; Supabase como backend-as-a-service (PostgreSQL, autenticación JWT, Row Level Security y Edge Functions); y despliegue serverless en Cloudflare Workers. Un aporte técnico relevante del proyecto fue la consolidación del dashboard financiero en su versión 2.0, mediante un estado global unificado (FinanceContext), eliminación de datos ficticios, normalización numérica multi-formato y validación automática de ecuaciones contables, garantizando trazabilidad entre el escáner OCR, la calculadora y los indicadores gerenciales.

Los resultados evidencian que la plataforma permite digitalizar estados de situación financiera y de resultados (P&L), validar la consistencia contable en tiempo real, personalizar la experiencia según el tipo de entidad (natural o jurídica) y el sector económico, y facilitar la toma de decisiones tributarias con apoyo de IA contextual. Se concluye que Contalogic AI+ constituye una propuesta viable, escalable y pertinente para democratizar herramientas financiero-tributarias avanzadas en el país, abriendo líneas futuras de integración con APIs de la DIAN, persistencia histórica de estados financieros y modelos predictivos.""",
    )

    kw = doc.add_paragraph()
    kw.paragraph_format.line_spacing = 1.5
    kw.paragraph_format.space_before = Pt(12)
    kr = kw.add_run(
        "Palabras clave: "
    )
    kr.bold = True
    kr.font.name = "Times New Roman"
    kr.font.size = Pt(12)
    kw.add_run(
        "gestión financiera, tributación colombiana, OCR, inteligencia artificial, "
        "Contalogic AI+, fintech, Supabase, React, cumplimiento DIAN."
    ).font.name = "Times New Roman"

    doc.add_page_break()

    # --- INTRODUCCIÓN ---
    doc.add_heading("INTRODUCCIÓN", level=1)
    add_paragraphs(
        doc,
        """La transformación digital ha redefinido la manera en que las organizaciones y los individuos administran su información financiera. Sin embargo, en el contexto colombiano persiste una brecha significativa entre las soluciones contables tradicionales —muchas veces costosas, rígidas o orientadas exclusivamente a empresas medianas y grandes— y las necesidades reales de microempresarios, profesionales independientes, pequeñas empresas y contadores que requieren herramientas integradas, accesibles y adaptadas a la normativa nacional.

En la práctica cotidiana, la gestión financiera y tributaria suele fragmentarse en múltiples artefactos: estados contables en PDF o imagen, registros manuales en Excel, calendarios de vencimientos en papel o aplicaciones genéricas, y consultas normativas realizadas de forma empírica o mediante buscadores no especializados. Esta dispersión dificulta la trazabilidad de la información, incrementa la probabilidad de errores en ecuaciones contables básicas —como Activos = Pasivos + Patrimonio— y limita la capacidad de anticipar obligaciones ante la DIAN o de identificar oportunidades legítimas de optimización tributaria.

Paralelamente, el avance de tecnologías como el reconocimiento óptico de caracteres (OCR) en el navegador y la inteligencia artificial generativa abre nuevas posibilidades para automatizar tareas que históricamente demandaban digitación manual y conocimiento especializado. No obstante, la mera disponibilidad tecnológica no garantiza por sí sola una solución pertinente: es necesario integrar dichas capacidades dentro de un diseño coherente, seguro y alineado al marco legal colombiano, con mecanismos de control de acceso, protección de datos personales y validación de la información financiera procesada.

En respuesta a este escenario se desarrolló Contalogic AI+ — Terminal Financiera Universal, una plataforma web modular concebida como un entorno unificado para el análisis financiero, la asistencia tributaria inteligente y el cumplimiento normativo. La aplicación reúne módulos de autenticación y perfiles personalizados, escáner OCR multi-formato, calculadora de ingresos y gastos por sector, dashboard con recordatorios DIAN, diccionario tributario, panel de ahorro de impuestos, asistente LEX AI y consola de administración, todo ello sobre una arquitectura moderna basada en React, Supabase y Cloudflare Workers.

El desarrollo del proyecto transcurrió de manera iterativa. La versión 1.0 (Beta) validó la viabilidad funcional de los módulos principales, pero evidenció problemas de consistencia derivados de fuentes de datos múltiples y valores ficticios hardcodeados. La versión 2.0 (Definitiva) abordó estas limitaciones mediante la implementación de un estado financiero global unificado, sincronización automática entre el escáner OCR y el dashboard, normalización numérica robusta y validación contable post-procesamiento, consolidando una única fuente de verdad para los indicadores financieros.

El presente documento corresponde a la monografía que sustenta académica y técnicamente este desarrollo. Su propósito es describir el problema investigado, justificar la pertinencia de la solución propuesta, presentar la arquitectura y funcionalidades de Contalogic AI+, documentar su evolución metodológica y aportar conclusiones sobre su utilidad en el ecosistema financiero-tributario colombiano. En consecuencia, el trabajo no se limita a enumerar tecnologías empleadas, sino que busca demostrar cómo la integración deliberada de OCR, analítica financiera e inteligencia artificial puede contribuir a mejorar la calidad, trazabilidad y oportunidad de la información contable y tributaria disponible para los usuarios.""",
    )

    doc.add_page_break()

    # --- PROBLEMA DE INVESTIGACIÓN ---
    doc.add_heading("DESCRIPCIÓN DEL PROBLEMA DE INVESTIGACIÓN", level=1)
    add_paragraphs(
        doc,
        """El problema de investigación que motiva este trabajo puede formularse de la siguiente manera: ¿cómo superar la fragmentación, inconsistencia y baja trazabilidad en la gestión financiera y tributaria de contribuyentes colombianos mediante una plataforma web integrada que combine extracción automatizada de estados contables, validación financiera en tiempo real y asistencia inteligente conforme al marco normativo nacional?

Para comprender la magnitud de este problema es necesario analizarlo desde tres dimensiones interrelacionadas: operativa, técnica y normativa.""",
    )

    doc.add_heading("Dimensión operativa", level=2)
    add_paragraphs(
        doc,
        """En el ámbito operativo, personas naturales y jurídicas enfrentan dificultades recurrentes para consolidar información financiera proveniente de distintas fuentes. Los estados de situación financiera y los estados de resultados (P&L) suelen recibirse en formatos no estructurados —PDF escaneados, fotografías, archivos Excel o documentos Word— lo que obliga a procesos manuales de transcripción, susceptibles de errores humanos y consumo excesivo de tiempo.

Asimismo, el registro cotidiano de ingresos y gastos muchas veces se realiza de forma separada del análisis de estados financieros formales, impidiendo una visión integral del desempeño económico. Los recordatorios de obligaciones tributarias ante la DIAN dependen con frecuencia de medios externos al sistema contable (calendarios físicos, notas, alertas genéricas), sin personalización según el tipo de contribuyente, el sector económico o el último dígito del NIT. Esta desarticulación genera riesgos de incumplimiento, sanciones y decisiones tardías en materia fiscal.""",
    )

    doc.add_heading("Dimensión técnica", level=2)
    add_paragraphs(
        doc,
        """Desde la perspectiva técnica, numerosas soluciones existentes abordan de manera parcial los retos descritos. Herramientas de OCR genéricas extraen texto, pero no interpretan partidas contables ni validan ecuaciones fundamentales. Software contable tradicional centraliza registros, pero muchas veces resulta costoso, complejo o poco adaptado a usuarios no especializados. Asistentes de inteligencia artificial ofrecen respuestas generales, pero sin contexto sobre la naturaleza del contribuyente (persona natural o jurídica), su sector económico o los datos financieros reales del usuario.

Durante el desarrollo inicial de Contalogic AI+ se identificó un problema técnico adicional: la coexistencia de múltiples fuentes de verdad dentro de la propia aplicación. El escáner OCR, la calculadora manual y el dashboard operaban con estados parcialmente desconectados, lo que producía inconsistencias entre totales extraídos, indicadores calculados y valores mostrados al usuario. También se detectó la presencia de datos ficticios hardcodeados en componentes de analítica, comprometiendo la confiabilidad del sistema. Estas fallas evidenciaron que la integración funcional no es un aspecto accesorio, sino un requisito central para cualquier plataforma que aspire a soportar decisiones financieras y tributarias.""",
    )

    doc.add_heading("Dimensión normativa", level=2)
    add_paragraphs(
        doc,
        """En la dimensión normativa, el entorno colombiano impone exigencias adicionales. El Estatuto Tributario y las disposiciones de la DIAN regulan obligaciones diferenciadas según el tipo de contribuyente, el régimen aplicable y la actividad económica. La interpretación de artículos específicos —por ejemplo, en materia de deducibilidad de gastos, retenciones, IVA, renta o ICA— requiere conocimiento especializado que no siempre está al alcance de pequeños empresarios o usuarios finales.

Por otro lado, cualquier plataforma que procese información financiera y personal debe observar la Ley 1581 de 2012 y el Decreto 1377 de 2013, garantizando confidencialidad, aislamiento de datos por usuario y mecanismos de autenticación seguros. En este sentido, el problema no consiste únicamente en automatizar cálculos, sino en hacerlo dentro de un marco de cumplimiento legal, responsabilidad y trazabilidad.""",
    )

    doc.add_heading("Planteamiento del problema", level=2)
    add_paragraphs(
        doc,
        """En síntesis, el problema de investigación se centra en la ausencia de una solución integral, accesible y contextualizada que permita a contribuyentes y profesionales contables en Colombia:

• Digitalizar estados financieros en múltiples formatos sin digitación manual extensiva.
• Unificar en una sola plataforma el registro operativo, el análisis de estados formales y los recordatorios tributarios.
• Validar automáticamente la consistencia contable de la información procesada.
• Recibir orientación tributaria personalizada con apoyo de inteligencia artificial.
• Acceder a reportes exportables (PDF/Excel) y tableros de control confiables.
• Operar bajo estándares de seguridad, roles de usuario y protección de datos personales.

La carencia de una herramienta que articule de manera coherente estos elementos perpetúa ineficiencias operativas, incrementa el riesgo de errores financieros y tributarios, y limita la capacidad de los usuarios para anticipar obligaciones y aprovechar oportunidades legítimas de ahorro fiscal. Contalogic AI+ se plantea como respuesta directa a este vacío, orientándose hacia la construcción de una terminal financiera universal que integre tecnología, contabilidad y normativa en un mismo entorno digital.""",
    )

    doc.add_page_break()

    # --- JUSTIFICACIÓN ---
    doc.add_heading("JUSTIFICACIÓN", level=1)
    add_paragraphs(
        doc,
        """La realización de este proyecto y la elaboración de la presente monografía se sustentan en cuatro tipos de justificación: académica, tecnológica, social y legal. Cada una de ellas refuerza la pertinencia de Contalogic AI+ como contribución al campo de la ingeniería de software aplicada al sector financiero-tributario colombiano.""",
    )

    doc.add_heading("Justificación académica", level=2)
    add_paragraphs(
        doc,
        """Desde el punto de vista académico, el desarrollo de Contalogic AI+ permite demostrar la aplicación integrada de conocimientos propios de la ingeniería de sistemas, la arquitectura de software, la programación web moderna, la gestión de bases de datos, la inteligencia artificial y el análisis financiero. El proyecto trasciende el ejercicio de codificación aislada al evidenciar un ciclo completo de investigación aplicada: identificación del problema, levantamiento de requerimientos, diseño arquitectónico, implementación iterativa, refactorización orientada a calidad de datos y documentación formal conforme al formato FSE-32.

Adicionalmente, la consolidación del dashboard financiero en la versión 2.0 constituye un caso de estudio relevante sobre gestión de estado en aplicaciones React, normalización de datos numéricos multi-formato y validación de reglas de negocio contables en tiempo de ejecución. Estos aprendizajes aportan valor pedagógico para futuros desarrollos académicos en fintech, sistemas de información contable y plataformas con procesamiento documental asistido por IA.""",
    )

    doc.add_heading("Justificación tecnológica", level=2)
    add_paragraphs(
        doc,
        """La justificación tecnológica reside en la oportunidad de aprovechar un stack moderno, accesible y escalable para construir una solución que antes requería infraestructura compleja y costosa. El uso de React 19 con TypeScript y TanStack Router/Start permite desarrollar una interfaz reactiva y mantenible; Supabase provee autenticación, persistencia relacional, Row Level Security y Edge Functions sin necesidad de administrar servidores propios; Cloudflare Workers habilita despliegue serverless de baja latencia; y bibliotecas como Tesseract.js, pdf.js, mammoth y xlsx posibilitan procesamiento OCR directamente en el navegador, reduciendo la exposición de documentos sensibles.

Contalogic AI+ demuestra que es viable integrar OCR, analítica financiera, exportación documental y asistentes de IA especializados en una sola plataforma coherente. La arquitectura implementada —con módulos de autenticación, escáner OCR, calculadora, dashboard, LEX AI, ahorro tributario, calendario DIAN, diccionario tributario y consola de administración— evidencia un diseño modular que facilita evolución futura, mantenimiento y escalamiento funcional.""",
    )

    doc.add_heading("Justificación social y económica", level=2)
    add_paragraphs(
        doc,
        """En el plano social y económico, la plataforma responde a una necesidad concreta de segmentos que históricamente han tenido acceso limitado a herramientas financiero-tributarias avanzadas: microempresarios, freelancers, pequeñas empresas, estudiantes de ciencias contables y profesionales que requieren agilizar procesos repetitivos sin sacrificar rigor técnico.

Al automatizar la extracción de datos contables, centralizar recordatorios DIAN personalizados y ofrecer recomendaciones de ahorro tributario basadas en información real del usuario, Contalogic AI+ contribuye a:

• Reducir tiempos de digitación y consolidación de estados financieros.
• Disminuir errores en ecuaciones contables básicas e indicadores derivados.
• Mejorar la oportunidad en el cumplimiento de obligaciones tributarias.
• Facilitar la comprensión de conceptos del Estatuto Tributario mediante LEX AI y el diccionario integrado.
• Democratizar el acceso a analítica financiera antes reservada a soluciones empresariales de mayor costo.

Estos beneficios tienen impacto directo en la productividad de contadores y empresarios, en la reducción de riesgos de sanciones por incumplimiento y en la toma de decisiones económicas más informadas, lo cual resulta especialmente relevante en un contexto de creciente formalización y digitalización de obligaciones ante la DIAN.""",
    )

    doc.add_heading("Justificación legal y ética", level=2)
    add_paragraphs(
        doc,
        """Finalmente, la justificación legal y ética se fundamenta en el compromiso del proyecto con el marco normativo colombiano. Contalogic AI+ fue diseñado considerando el Estatuto Tributario, las disposiciones de la DIAN y la Ley 1581 de 2012 sobre protección de datos personales. La implementación de autenticación JWT, políticas de Row Level Security por usuario, control de roles (super_admin, admin, user) y procesamiento OCR en el cliente refleja una preocupación explícita por la confidencialidad, integridad y trazabilidad de la información financiera.

Es importante señalar que la plataforma no sustituye el criterio profesional de un contador público ni constituye asesoría jurídica vinculante; no obstante, sí actúa como herramienta de apoyo que orienta, organiza y facilita el acceso a información relevante, siempre dentro de un marco de uso responsable de la inteligencia artificial. Esta distinción refuerza la pertinencia ética del proyecto en un dominio regulado donde la precisión, la transparencia y la protección de datos son requisitos ineludibles.""",
    )

    doc.add_heading("Conclusión de la justificación", level=2)
    add_paragraphs(
        doc,
        """En conjunto, las justificaciones expuestas demuestran que Contalogic AI+ no responde a un interés meramente tecnológico, sino a una necesidad real, verificable y contextualizada. La plataforma aporta valor académico como artefacto de investigación aplicada, valor tecnológico como integración coherente de OCR, IA y analítica financiera, valor social como herramienta de inclusión digital en materia tributaria, y valor legal como desarrollo alineado al marco normativo colombiano. Por estas razones, su estudio, documentación y evaluación mediante la presente monografía resultan plenamente pertinentes y necesarios.""",
    )

    doc.save(OUT)
    print(f"Documento generado: {OUT}")


if __name__ == "__main__":
    main()
